"""
Agentic MOC -> Equipment mapper (LangGraph + Claude on Azure AI Foundry).

Single self-contained tool: reads the MOC list from the FIRST sheet of the
workbook only (the reviewed sheet with the "Remark 10_Sep" column), keeps
only rows for one location (default BOKOR) whose remark starts with
"INCLUDE", and for every remaining MOC asks a Claude agent to decide which
equipment record (if any) from the plant's equipment master the MOC title is
actually about.

The agent is a small LangGraph state machine:

    START -> retrieve -> decide --(confident)--> END
                ^            |
                |     (not confident, attempts left)
                |            v
                +-------- broaden

- retrieve: deterministic tag/keyword search that narrows ~24,000 equipment
  rows down to a short candidate list (fast, no LLM call - equipment master
  is far too large to hand to an LLM in one go).
- decide: Claude reads the MOC title against the numbered candidate list and
  returns a structured decision (match / no match, which one, confidence,
  reasoning). Claude never invents an equipment number outside the list.
- broaden: if Claude isn't confident, the candidate list is widened (3x,
  up to a cap) and the agent tries again, up to MAX_ATTEMPTS times.

Install:
    pip install langgraph anthropic python-dotenv pydantic openpyxl

Run:
    python moc_equipment_agent.py                 # all Bokor MOCs
    python moc_equipment_agent.py --concurrency 10 # process 10 MOCs' Claude calls in parallel
    python moc_equipment_agent.py --limit 10       # quick test on first 10 MOCs
    python moc_equipment_agent.py --dry-run        # skip Claude calls, just show candidates
    python moc_equipment_agent.py --location BETTY --moc-file "other.xlsx"
"""

import argparse
import json
import math
import os
import re
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from typing import Any, Dict, List, Literal, Optional, TypedDict

from dotenv import load_dotenv
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field, ValidationError

from anthropic import AnthropicFoundry
from langgraph.graph import END, START, StateGraph

load_dotenv()

# --------------------------------------------------------------------------
# Configuration (can be overridden via command-line arguments)
# --------------------------------------------------------------------------

DEFAULT_MOC_FILE = "EMOCV2_List_30_3_2026 (BOBELA)_review 10 Sept.xlsx"
DEFAULT_EQUIPMENT_FILE = "Equipment_Bokor.xlsx"
DEFAULT_EQUIPMENT_SHEET = "Sheet1"
DEFAULT_LOCATION = "BOKOR"
DEFAULT_OUTPUT_FILE = "MOC_to_Equipment_Mapping.xlsx"
DEFAULT_REMARK_FILTER = "INCLUDE"  # only rows whose "Remark 10_Sep" starts with this are processed

INITIAL_K = 12       # candidates shown to Claude on the first attempt
MAX_ATTEMPTS = 3     # how many times the graph will broaden and retry
MAX_K = 100          # candidate-list size cap

STOPWORDS = {
    "WITH", "WITHOUT", "FROM", "INTO", "ONTO", "ABOVE", "BELOW", "ABSENT",
    "PROLONG", "PROLONGED", "CONTINUE", "CONTINUING", "OPERATE", "OPERATING",
    "OPERATION", "FLOWING", "PERMANENT", "PERMANENTLY", "ESTABLISH",
    "FINDINGS", "GUIDELINE", "ASSURANCE", "INTEGRITY", "ALLOWABLE",
    "THICKNESS", "RECORDS", "WELDING", "LIMIT", "SYSTEM", "UNAVAILABLE",
    "STANDARD", "PHILOSOPHY", "DISCONTINUE", "DISCONTINUED", "INSTALLATION",
    "SWITCHOVER", "CHEMICAL", "DEVICES", "INDIVIDUAL", "CONTROL", "RATE",
    "HAZID", "WHILE", "THAT", "THIS", "THEIR", "THERE", "WHICH", "UNTIL",
    "DURING", "AFTER", "BEFORE", "BEING", "HAVE", "HAS", "FAULTY",
}

# --------------------------------------------------------------------------
# Text helpers
# --------------------------------------------------------------------------


def compact(value):
    """Uppercase and strip everything but letters/digits."""
    return re.sub(r"[^A-Z0-9]", "", str(value or "").upper())


def has_letter_and_digit(s):
    return any(c.isalpha() for c in s) and any(c.isdigit() for c in s)


TOKEN_RE = re.compile(r"\b[A-Z0-9]+(?:[-/][A-Z0-9]+)*\b")


def expand_slash_token(token):
    """Expand compound tags like 'P4210A/B/C' or 'S-5420A/B' into full tags."""
    if "/" not in token:
        return [token]
    parts = token.split("/")
    first = parts[0]
    # Prefix must end in a digit; the varying part is the trailing 1-2 letters
    # (e.g. "S-5420A" -> prefix "S-5420", suffix "A").
    prefix_match = re.match(r"^(.*\d)([A-Z]{1,2})$", first)
    prefix = prefix_match.group(1) if prefix_match else first
    expanded = [first]
    for part in parts[1:]:
        if len(part) <= 2 and part.isalnum():
            expanded.append(prefix + part)
        else:
            expanded.append(part)
    return expanded


def extract_tag_candidates(title):
    """Pull candidate equipment-tag strings out of an MOC title."""
    if not title:
        return []
    upper_title = str(title).upper()
    candidates = set()
    for raw_token in TOKEN_RE.findall(upper_title):
        for token in expand_slash_token(raw_token):
            full = compact(token)
            if len(full) >= 3 and has_letter_and_digit(full):
                candidates.add(full)
            if "-" in token:
                for part in token.split("-"):
                    part_c = compact(part)
                    if len(part_c) >= 3 and has_letter_and_digit(part_c):
                        candidates.add(part_c)
    return sorted(candidates)


def significant_words(text):
    words = re.findall(r"[A-Z]{4,}", str(text or "").upper())
    return {w for w in words if w not in STOPWORDS}


def find_column(headers, *names):
    """Find a column index (0-based) whose header matches one of `names`."""
    normalized = {str(h).strip().lower() if h else "": i for i, h in enumerate(headers)}
    for name in names:
        idx = normalized.get(name.strip().lower())
        if idx is not None:
            return idx
    raise KeyError(f"None of the expected column headers {names} found in {headers}")


# --------------------------------------------------------------------------
# Data loading
# --------------------------------------------------------------------------


def load_mocs(moc_file, location, remark_filter=DEFAULT_REMARK_FILTER):
    """Load MOCs from the workbook's FIRST sheet only (by position, regardless
    of its name), keeping only rows for `location` whose "Remark 10_Sep"
    column starts with `remark_filter` (case-insensitive)."""
    wb = load_workbook(moc_file, data_only=True)
    ws = wb.worksheets[0]
    print(f"  (using first sheet: '{ws.title}')")
    rows = list(ws.iter_rows(values_only=True))
    headers = rows[0]

    col_opu = find_column(headers, "OPU")
    col_loc8 = find_column(headers, "Location 8")
    col_emoc_no = find_column(headers, "EMOC No")
    col_remark = find_column(headers, "Remark 10_Sep", "Remark")
    col_title = find_column(headers, "Title")
    col_nature = find_column(headers, "Nature of Change")
    col_workflow = find_column(headers, "Workflow Status")
    col_status = find_column(headers, "MOC Status")
    col_initiate = find_column(headers, "Initiate Date")

    mocs = []
    for row in rows[1:]:
        if row is None or all(v is None for v in row):
            continue
        loc8 = row[col_loc8] if col_loc8 < len(row) else None
        if not loc8 or location.upper() not in str(loc8).upper():
            continue
        remark = row[col_remark] if col_remark < len(row) else None
        if not remark or not str(remark).strip().upper().startswith(remark_filter.upper()):
            continue
        mocs.append(
            {
                "opu": row[col_opu] if col_opu < len(row) else None,
                "location_8": loc8,
                "emoc_no": row[col_emoc_no] if col_emoc_no < len(row) else None,
                "remark": remark,
                "title": row[col_title] if col_title < len(row) else None,
                "nature_of_change": row[col_nature] if col_nature < len(row) else None,
                "workflow_status": row[col_workflow] if col_workflow < len(row) else None,
                "moc_status": row[col_status] if col_status < len(row) else None,
                "initiate_date": row[col_initiate] if col_initiate < len(row) else None,
            }
        )
    return mocs


def load_equipment(equipment_file, sheet_name, location=None):
    wb = load_workbook(equipment_file, data_only=True)
    ws = wb[sheet_name]
    rows = ws.iter_rows(values_only=True)
    headers = next(rows)

    col_equipment = find_column(headers, "Equipment")
    col_obj_desc = find_column(headers, "Object Description")
    col_func_loc = find_column(headers, "Functional Location")
    col_func_desc = find_column(headers, "FunctLocDescrip.", "FunctLocDescrip")
    col_location = find_column(headers, "Location")
    col_tech_ident = find_column(headers, "Technical IdentNo.", "Technical IdentNo")
    col_mat_desc = find_column(headers, "Material Description")

    equipment = []
    for row in rows:
        if row is None or all(v is None for v in row):
            continue
        loc = row[col_location] if col_location < len(row) else None
        if location and (not loc or location.upper() not in str(loc).upper()):
            continue

        obj_desc = row[col_obj_desc] if col_obj_desc < len(row) else None
        func_loc = row[col_func_loc] if col_func_loc < len(row) else None
        func_desc = row[col_func_desc] if col_func_desc < len(row) else None
        tech_ident = row[col_tech_ident] if col_tech_ident < len(row) else None
        mat_desc = row[col_mat_desc] if col_mat_desc < len(row) else None

        equipment.append(
            {
                "equipment": row[col_equipment] if col_equipment < len(row) else None,
                "object_description": obj_desc,
                "functional_location": func_loc,
                "func_loc_description": func_desc,
                "location": loc,
                "technical_ident_no": tech_ident,
                "material_description": mat_desc,
                "_c_tech_ident": compact(tech_ident),
                "_c_func_loc": compact(func_loc),
                "_c_func_desc": compact(func_desc),
                "_c_obj_desc": compact(obj_desc),
                "_c_mat_desc": compact(mat_desc),
                "_words": significant_words(obj_desc) | significant_words(func_desc),
            }
        )
    return equipment


# --------------------------------------------------------------------------
# Deterministic candidate retrieval (this is the "tool" the agent calls to
# narrow the search space - it never decides the final answer by itself)
# --------------------------------------------------------------------------


def tag_match_score(tag, eq):
    """Score how well a single compact tag matches one equipment record."""
    if tag and tag == eq["_c_tech_ident"]:
        return 100
    if tag and tag in eq["_c_tech_ident"]:
        return 90
    if tag and tag in eq["_c_func_loc"]:
        return 80
    if tag and tag in eq["_c_func_desc"]:
        return 60
    if tag and tag in eq["_c_obj_desc"]:
        return 50
    if tag and tag in eq["_c_mat_desc"]:
        return 40
    return 0


def match_by_tags(candidates, equipment):
    """Return {equipment_index: (score, matched_tag)} for tag-based hits."""
    best = {}
    for tag in candidates:
        for idx, eq in enumerate(equipment):
            score = tag_match_score(tag, eq)
            if score and (idx not in best or score > best[idx][0]):
                best[idx] = (score, tag)
    return best


def build_word_doc_frequency(equipment):
    freq = defaultdict(int)
    for eq in equipment:
        for w in eq["_words"]:
            freq[w] += 1
    return freq


def match_by_fuzzy_words(title, equipment, word_doc_freq, n_docs):
    words = significant_words(title)
    if not words:
        return {}
    scores = {}
    for idx, eq in enumerate(equipment):
        overlap = words & eq["_words"]
        if not overlap:
            continue
        score = sum(math.log(n_docs / word_doc_freq[w]) for w in overlap)
        if score > 0:
            scores[idx] = (min(35, 8 * score), "TEXT:" + ",".join(sorted(overlap)))
    return scores


def get_candidates(moc, equipment, word_doc_freq, n_docs, k):
    """Narrow the equipment master down to at most `k` candidates for a MOC."""
    tags = extract_tag_candidates(moc["title"])
    tag_hits = match_by_tags(tags, equipment)
    ranked = sorted(tag_hits.items(), key=lambda kv: kv[1][0], reverse=True)
    ordered = [(idx, score, tag) for idx, (score, tag) in ranked]

    if len(ordered) < k:
        fuzzy_hits = match_by_fuzzy_words(moc["title"], equipment, word_doc_freq, n_docs)
        have = {idx for idx, _, _ in ordered}
        for idx, (score, tag) in sorted(fuzzy_hits.items(), key=lambda kv: kv[1][0], reverse=True):
            if idx in have:
                continue
            ordered.append((idx, score, tag))
            have.add(idx)
            if len(ordered) >= k:
                break

    ordered = ordered[:k]
    return tags, [
        {
            "idx": idx,
            "heuristic_score": score,
            "heuristic_tag": tag,
            "equipment": equipment[idx]["equipment"],
            "object_description": equipment[idx]["object_description"],
            "functional_location": equipment[idx]["functional_location"],
            "func_loc_description": equipment[idx]["func_loc_description"],
            "technical_ident_no": equipment[idx]["technical_ident_no"],
        }
        for idx, score, tag in ordered
    ]


# --------------------------------------------------------------------------
# Claude client (Azure AI Foundry deployment)
# --------------------------------------------------------------------------


class ClaudeFoundryClient:
    """Claude deployment on Microsoft/Azure AI Foundry, via the native
    Anthropic Messages API (https://{resource}.services.ai.azure.com/anthropic)."""

    def __init__(self, resource: str, deployment: str, api_key: str):
        self.client = AnthropicFoundry(api_key=api_key, resource=resource)
        self.deployment = deployment

    def chat(self, system_prompt: str, user_prompt: str, max_tokens: int = 600) -> str:
        response = self.client.messages.create(
            model=self.deployment,
            max_tokens=max_tokens,
            system=system_prompt,
            messages=[{"role": "user", "content": user_prompt}],
        )
        return "".join(block.text for block in response.content if block.type == "text")


def get_claude_client() -> ClaudeFoundryClient:
    return ClaudeFoundryClient(
        resource=os.environ["AZURE_CLAUDE_RESOURCE"],
        deployment=os.environ["AZURE_CLAUDE_DEPLOYMENT"],
        api_key=os.environ["AZURE_CLAUDE_API_KEY"],
    )


# --------------------------------------------------------------------------
# Structured decision schema + prompt
# --------------------------------------------------------------------------


class EquipmentDecision(BaseModel):
    is_match: bool = Field(description="True only if one candidate is genuinely the equipment this MOC is about")
    selected_index: Optional[int] = Field(default=None, description="1-based index into the candidate list, or null")
    confidence: Literal["high", "medium", "low"] = Field(
        description="How strong the MATCH EVIDENCE is (only meaningful when is_match is true)"
    )
    matched_reason: str = Field(description="The tag/keyword or phrase that ties the title to the equipment")
    rationale: str = Field(description="One or two sentence explanation")


SYSTEM_PROMPT = """You are a process/reliability engineer reviewing Management of Change (MOC) \
records for an oil & gas facility. You are given an MOC title (and related fields) plus a \
numbered short-list of candidate equipment pulled from the plant's equipment master.

Decide whether exactly ONE candidate is genuinely the equipment the MOC is about.

Rules of thumb:
- Equipment tags embedded in the title (e.g. "K-2510", "P-4110A", "FD-067", "685FT-001") should \
match an equipment's Technical IdentNo. or Functional Location, not just share a coincidental \
number.
- Well names / pipeline names (e.g. "BO-208S", "PL69") are not process equipment - only match \
them if a candidate is clearly the actual wellhead/pipeline equipment record itself.
- Site-wide, procedural, or documentation-only MOCs (e.g. about a whole platform, missing \
records, a report, a philosophy) usually have NO correct single-equipment match. In that case \
set is_match to false.
- Never invent an equipment number that is not in the candidate list. If none of the candidates \
fit, set is_match to false and selected_index to null.

IMPORTANT - "confidence" means the strength of the MATCH EVIDENCE, not how sure you feel in \
general. Calibrate it only for the equipment you selected (leave it "low" whenever is_match is \
false - a confident "no match" is not a confident match):
- "high": the title's tag/identifier exactly matches this equipment's Technical IdentNo. or \
Functional Location (e.g. title says "K-2510", equipment's Technical IdentNo. is "BOCP-K2510").
- "medium": the equipment is clearly the right general item (same tag family, same area/system) \
but the match relies on partial or contextual evidence rather than an exact identifier hit.
- "low": you selected this candidate as the closest available option, but the evidence is \
thin/inferred (e.g. only a text/keyword overlap, no tag confirmation) and it should be \
manually verified.

Respond with ONLY strict JSON (no markdown fences, no prose outside the JSON) matching exactly:
{"is_match": true|false, "selected_index": <int or null>, "confidence": "high"|"medium"|"low", \
"matched_reason": "<string>", "rationale": "<string>"}
"""


def build_user_prompt(moc: Dict[str, Any], candidates: List[Dict[str, Any]]) -> str:
    lines = [
        f"MOC No: {moc.get('emoc_no')}",
        f"Title: {moc.get('title')}",
        f"Nature of Change: {moc.get('nature_of_change')}",
        f"OPU: {moc.get('opu')}",
        f"Location: {moc.get('location_8')}",
        "",
        "Candidate equipment:",
    ]
    for i, c in enumerate(candidates, start=1):
        lines.append(
            f"[{i}] Equipment No: {c.get('equipment')} | "
            f"Object Description: {c.get('object_description')} | "
            f"Functional Location: {c.get('functional_location')} | "
            f"FunctLocDescrip.: {c.get('func_loc_description')} | "
            f"Technical IdentNo.: {c.get('technical_ident_no')}"
        )
    if not candidates:
        lines.append("(no candidates found)")
    return "\n".join(lines)


def strip_code_fences(text: str) -> str:
    text = text.strip()
    text = re.sub(r"^```(?:json)?", "", text).strip()
    text = re.sub(r"```$", "", text).strip()
    return text


def ask_claude_to_match(client: ClaudeFoundryClient, moc: Dict[str, Any], candidates: List[Dict[str, Any]]) -> EquipmentDecision:
    user_prompt = build_user_prompt(moc, candidates)
    last_error = None
    for _ in range(2):
        try:
            raw = client.chat(SYSTEM_PROMPT, user_prompt)
            raw = strip_code_fences(raw)
            data = json.loads(raw)
            decision = EquipmentDecision.model_validate(data)
            if decision.selected_index is not None and not (1 <= decision.selected_index <= len(candidates)):
                decision.selected_index = None
                decision.is_match = False
            return decision
        except (json.JSONDecodeError, ValidationError) as exc:
            last_error = exc
            user_prompt += (
                f"\n\nYour previous reply could not be parsed as valid JSON matching the schema "
                f"({exc}). Reply again with ONLY the corrected JSON object."
            )
    return EquipmentDecision(
        is_match=False,
        selected_index=None,
        confidence="low",
        matched_reason="",
        rationale=f"Claude response could not be parsed after retries: {last_error}",
    )


# --------------------------------------------------------------------------
# LangGraph agent: retrieve -> decide -> (broaden -> retrieve -> decide)*
# --------------------------------------------------------------------------


class AgentState(TypedDict):
    moc: Dict[str, Any]
    k: int
    attempt: int
    extracted_tags: List[str]
    candidates: List[Dict[str, Any]]
    decision: Optional[Dict[str, Any]]


def make_graph(client: Optional[ClaudeFoundryClient], equipment, word_doc_freq, n_docs, dry_run: bool):
    def node_retrieve(state: AgentState) -> AgentState:
        tags, candidates = get_candidates(state["moc"], equipment, word_doc_freq, n_docs, state["k"])
        return {**state, "extracted_tags": tags, "candidates": candidates}

    def node_decide(state: AgentState) -> AgentState:
        if dry_run or client is None:
            decision = EquipmentDecision(
                is_match=bool(state["candidates"]),
                selected_index=1 if state["candidates"] else None,
                confidence="low",
                matched_reason="dry-run (no LLM call made)",
                rationale="dry-run",
            )
        else:
            decision = ask_claude_to_match(client, state["moc"], state["candidates"])
        return {**state, "decision": decision.model_dump()}

    def route_after_decide(state: AgentState) -> str:
        decision = state["decision"]
        if decision["is_match"] and decision["confidence"] in ("high", "medium"):
            return "end"
        if state["attempt"] >= MAX_ATTEMPTS or state["k"] >= min(MAX_K, len(equipment)):
            return "end"
        return "broaden"

    def node_broaden(state: AgentState) -> AgentState:
        new_k = min(state["k"] * 3, MAX_K, len(equipment))
        return {**state, "k": new_k, "attempt": state["attempt"] + 1}

    workflow = StateGraph(AgentState)
    workflow.add_node("retrieve", node_retrieve)
    workflow.add_node("decide", node_decide)
    workflow.add_node("broaden", node_broaden)
    workflow.add_edge(START, "retrieve")
    workflow.add_edge("retrieve", "decide")
    workflow.add_conditional_edges("decide", route_after_decide, {"end": END, "broaden": "broaden"})
    workflow.add_edge("broaden", "retrieve")
    return workflow.compile()


# --------------------------------------------------------------------------
# Output
# --------------------------------------------------------------------------


def write_output(rows: List[Dict[str, Any]], output_file: str):
    wb = Workbook()
    ws = wb.active
    ws.title = "MOC_Equipment_Mapping"
    headers = [
        "EMOC No", "Title", "OPU", "Location 8", "Remark 10_Sep", "MOC Status", "Extracted Tags",
        "Confidence", "Selected Equipment No", "Object Description",
        "Functional Location", "FunctLocDescrip.", "Technical IdentNo.",
        "Matched Reason", "Rationale",
    ]
    ws.append(headers)
    for r in rows:
        ws.append(
            [
                r["moc"]["emoc_no"],
                r["moc"]["title"],
                r["moc"]["opu"],
                r["moc"]["location_8"],
                r["moc"]["remark"],
                r["moc"]["moc_status"],
                ", ".join(r["tags"]),
                r["decision"]["confidence"],
                r["equipment"]["equipment"] if r["equipment"] else None,
                r["equipment"]["object_description"] if r["equipment"] else None,
                r["equipment"]["functional_location"] if r["equipment"] else None,
                r["equipment"]["func_loc_description"] if r["equipment"] else None,
                r["equipment"]["technical_ident_no"] if r["equipment"] else None,
                r["decision"]["matched_reason"],
                r["decision"]["rationale"],
            ]
        )

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
        ws.column_dimensions[get_column_letter(col_idx)].width = max(14, min(50, len(header) + 4))
    ws.freeze_panes = "A2"
    wb.save(output_file)


# --------------------------------------------------------------------------
# Main
# --------------------------------------------------------------------------


def main():
    parser = argparse.ArgumentParser(description="Agentic (LangGraph + Claude on Azure Foundry) MOC-to-equipment mapper.")
    parser.add_argument("--moc-file", default=DEFAULT_MOC_FILE)
    parser.add_argument("--equipment-file", default=DEFAULT_EQUIPMENT_FILE)
    parser.add_argument("--equipment-sheet", default=DEFAULT_EQUIPMENT_SHEET)
    parser.add_argument("--location", default=DEFAULT_LOCATION)
    parser.add_argument("--output-file", default=DEFAULT_OUTPUT_FILE)
    parser.add_argument("--remark-filter", default=DEFAULT_REMARK_FILTER,
                         help='Only process MOCs whose "Remark 10_Sep" column starts with this (default: INCLUDE).')
    parser.add_argument("--limit", type=int, default=None, help="Only process the first N MOCs (testing).")
    parser.add_argument("--dry-run", action="store_true", help="Skip Claude calls; just show retrieved candidates.")
    parser.add_argument("--concurrency", type=int, default=5, help="Number of MOCs processed in parallel (Claude calls in flight at once).")
    args = parser.parse_args()

    print(f"Loading MOCs from '{args.moc_file}' (first sheet only) for location '{args.location}', "
          f"remark starting with '{args.remark_filter}'...")
    mocs = load_mocs(args.moc_file, args.location, args.remark_filter)
    if args.limit:
        mocs = mocs[: args.limit]
    print(f"  -> {len(mocs)} MOC rows to process.")

    print(f"Loading equipment from '{args.equipment_file}' [{args.equipment_sheet}]...")
    equipment = load_equipment(args.equipment_file, args.equipment_sheet, args.location)
    print(f"  -> {len(equipment)} equipment rows loaded.")

    word_doc_freq = build_word_doc_frequency(equipment)
    n_docs = len(equipment)

    client = None
    if not args.dry_run:
        print("Connecting to Claude deployment on Azure AI Foundry...")
        client = get_claude_client()

    app = make_graph(client, equipment, word_doc_freq, n_docs, args.dry_run)

    def process_one(moc: Dict[str, Any]) -> Dict[str, Any]:
        state: AgentState = {
            "moc": moc,
            "k": INITIAL_K,
            "attempt": 0,
            "extracted_tags": [],
            "candidates": [],
            "decision": None,
        }
        try:
            final_state = app.invoke(state)
        except Exception as exc:  # network/API errors shouldn't kill the whole run
            final_state = {
                **state,
                "candidates": [],
                "decision": EquipmentDecision(
                    is_match=False, selected_index=None, confidence="low",
                    matched_reason="", rationale=f"error: {exc}",
                ).model_dump(),
            }

        decision = final_state["decision"]
        candidates = final_state["candidates"]
        equipment_row = None
        if decision["is_match"] and decision["selected_index"]:
            equipment_row = candidates[decision["selected_index"] - 1]

        return {
            "moc": moc,
            "tags": final_state["extracted_tags"],
            "candidates": candidates,
            "attempts": final_state["attempt"],
            "decision": decision,
            "equipment": equipment_row,
        }

    print(f"Processing {len(mocs)} MOCs with up to {args.concurrency} in parallel...")
    results_by_index: Dict[int, Dict[str, Any]] = {}
    completed = 0
    with ThreadPoolExecutor(max_workers=args.concurrency) as executor:
        future_to_index = {executor.submit(process_one, moc): i for i, moc in enumerate(mocs)}
        for future in as_completed(future_to_index):
            index = future_to_index[future]
            result = future.result()
            results_by_index[index] = result
            completed += 1
            decision = result["decision"]
            print(f"  [{completed}/{len(mocs)}] {result['moc'].get('emoc_no')} -> "
                  f"{'MATCH' if decision['is_match'] else 'no match'} "
                  f"({decision['confidence']}, {result['attempts']} broaden attempt(s))")

    results = [results_by_index[i] for i in range(len(mocs))]
    matched_results = [r for r in results if r["decision"]["is_match"] and r["equipment"]]
    write_output(matched_results, args.output_file)
    print(f"\nDone. {len(matched_results)}/{len(results)} MOCs matched. "
          f"Written to '{args.output_file}' (unmatched MOCs are excluded from the file).")


if __name__ == "__main__":
    main()
