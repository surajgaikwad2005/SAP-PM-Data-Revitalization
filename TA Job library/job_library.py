"""Build the SD&TA Job Library for Bokor 2024.

Pipeline:
  1. Classify every worklist equipment into an Equipment Class (Level 2).
  2. For each class, map its equipments to Tasklist_2024 to get their task lines.
  3. Agent clusters equipments that share the SAME maintenance procedure into
     Group Counters, generalising the operation descriptions (strip equipment
     tags) and preserving task sequence.
  4. Write 3 sheets:
       Sheet 1: Equipment -> Class mapping
       Sheet 2: Job Library (TaskListGroup, Counter, Header, Plant, Operation,
                Duration, Unit) with a total-duration row per counter
       Sheet 3: TaskListGroup, Counter -> Equipments

LLM: Claude Sonnet via Azure Foundry, all calls traced through Langfuse
(agent name = "SD&TA Agent", project = "SAP PM Data Revitalization").
"""

import json
import os
import re
from pathlib import Path

import pandas as pd
from anthropic import Anthropic
from dotenv import load_dotenv
from langfuse import Langfuse, observe
from openpyxl.styles import Font, PatternFill

# ---------------------------------------------------------------------------
# Config
# ---------------------------------------------------------------------------
load_dotenv(Path(__file__).resolve().parent / ".env")

FOUNDRY_BASE_URL = "https://tridiagonal-ai-eastus2-resource.services.ai.azure.com/anthropic"
MODEL = "claude-sonnet-4-6"
AGENT_NAME = "SD&TA Agent"
LF_PROJECT = "SAP PM Data Revitalization"

BASE = Path(__file__).resolve().parents[1]
WORKLIST = BASE / "Processed Data" / "Cleanly_Mapped_Worklist_Bokor_2024.xlsx"
EQ_MASTER = BASE / "Processed Data" / "Equipment_Master_Bokor_2024.xlsx"
EQ_CLASS = BASE / "Raw data" / "Equipment Class.xlsx"
TASKLIST = BASE / "Merged Sheet" / "Tasklist_2024.xlsx"
OUTPUT = BASE / "Processed Data" / "Job_Library_Bokor_2024.xlsx"

PLANT = "M014"
PLANT_CODE = "SK"
LEVEL = 2  # equipment class level

anthropic_client = Anthropic(api_key=os.getenv("AZURE_AI_KEY"), base_url=FOUNDRY_BASE_URL)
langfuse = Langfuse()  # reads LANGFUSE_* env vars


# ---------------------------------------------------------------------------
# Equipment class resolver (level 2 preferred, walk up if needed)
# ---------------------------------------------------------------------------
def load_class_lookup():
    df = pd.read_excel(EQ_CLASS)[["Level", "Class", "Class Description"]].dropna(subset=["Class"])
    df.columns = ["Level", "Class", "Description"]
    df["Class"] = df["Class"].astype(str).str.strip().str.upper()
    return {
        r.Class: {"level": int(r.Level), "description": r.Description}
        for r in df.itertuples(index=False)
    }


def resolve_class(cp, lookup):
    """`1EJCL` -> walk prefixes, prefer Level 2 (or coarser)."""
    if pd.isna(cp):
        return None, None, None
    core = re.sub(r"^\d+", "", str(cp).strip().upper())
    for n in range(len(core), 0, -1):
        hit = lookup.get(core[:n])
        if hit and hit["level"] <= LEVEL:
            return core[:n], hit["level"], hit["description"]
    return None, None, None


# ---------------------------------------------------------------------------
# LLM call (traced in Langfuse)
# ---------------------------------------------------------------------------
FEW_SHOT = """EXAMPLE INPUT:
{
  "equipment_class": "MRP - PUMP",
  "equipments": [
    {"id": 1, "equipment_no": "P-101",
     "tasks": ["Isolate pump P-101", "Remove head of pump P-101",
               "Grease coupling of P-101", "Reassemble P-101"],
     "durations_hrs": [1.0, 2.0, 0.5, 1.5]},
    {"id": 2, "equipment_no": "P-202",
     "tasks": ["Isolate pump P-202", "Remove head of pump P-202",
               "Grease coupling of P-202", "Reassemble P-202"],
     "durations_hrs": [1.0, 2.0, 0.5, 1.5]},
    {"id": 3, "equipment_no": "P-305",
     "tasks": ["Megger test on motor of P-305", "Insulation resistance check P-305"],
     "durations_hrs": [1.0, 0.5]}
  ]
}

EXAMPLE OUTPUT:
{
  "counters": [
    {
      "header": "Pump Head Removal and Coupling Service",
      "operations": [
        {"description": "Isolate pump", "duration_hrs": 1.0},
        {"description": "Remove head of pump", "duration_hrs": 2.0},
        {"description": "Grease coupling", "duration_hrs": 0.5},
        {"description": "Reassemble pump", "duration_hrs": 1.5}
      ],
      "equipment_ids": [1, 2]
    },
    {
      "header": "Motor Insulation Testing",
      "operations": [
        {"description": "Megger test on motor", "duration_hrs": 1.0},
        {"description": "Insulation resistance check", "duration_hrs": 0.5}
      ],
      "equipment_ids": [3]
    }
  ]
}"""

SYSTEM = f"""You are an SAP PM maintenance planner building a Job Library for turnaround/shutdown activities on a Petronas offshore platform.

INPUT: For ONE equipment class, you receive a list of equipments. Each equipment has its
sequence of maintenance task steps and the per-step duration (hours). Order matters -
these are shutdown steps performed in that sequence.

JOB: Cluster equipments that share the SAME maintenance procedure (the same set of steps
in the same intent, even when wording differs because of equipment tags/numbers). Each
cluster becomes one "Group Counter". For each counter, produce the GENERALIZED operation
list - strip equipment numbers / tags, keep the procedure in neutral terms - preserving
the original sequence. Average per-step duration across the equipments in that counter.

RULES:
- Every input equipment_id MUST appear in exactly one counter.
- "header" <= 60 chars, title case, canonical name of the procedure.
- Operations are ordered as the steps occur in the task sequence.
- Deduplicate near-identical steps within a counter (e.g. two phrasings of "visual inspection").
- If an equipment is a one-off (unique procedure), it is a counter of its own.

Return STRICT JSON only, matching the example schema exactly.

{FEW_SHOT}"""


@observe(as_type="generation", name="cluster_tasks")
def call_llm(class_code, class_desc, payload):
    user_text = json.dumps(payload, ensure_ascii=False, default=str)
    resp = anthropic_client.messages.create(
        model=MODEL,
        max_tokens=4000,
        system=SYSTEM,
        messages=[{"role": "user", "content": user_text}],
    )
    text = "".join(b.text for b in resp.content if getattr(b, "type", None) == "text")

    langfuse.update_current_generation(
        model=MODEL,
        input=[{"role": "system", "content": SYSTEM},
               {"role": "user", "content": user_text}],
        output=text,
        usage_details={"input": resp.usage.input_tokens, "output": resp.usage.output_tokens},
        metadata={"agent": AGENT_NAME, "project": LF_PROJECT,
                  "class": class_code, "class_description": class_desc},
    )

    m = re.search(r"\{.*\}", text, re.DOTALL)
    return json.loads(m.group(0)) if m else {"counters": []}


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
@observe(name="build_job_library")
def build():
    langfuse.update_current_trace(name=f"{AGENT_NAME} - Build Job Library",
                                  metadata={"project": LF_PROJECT, "year": 2024})

    # ----- Load inputs
    wl = pd.read_excel(WORKLIST)
    em = pd.read_excel(EQ_MASTER)[["EquipmentNo", "ObjectDescription", "CatalogProfile"]]
    tl = pd.read_excel(TASKLIST)[["WorklistID", "EquipmentNo", "TaskNo", "Task", "Dur (Hrs)"]]
    lookup = load_class_lookup()

    # ----- Attach class to every worklist equipment
    wl["Suggested EquipmentNo"] = wl["Suggested EquipmentNo"].astype(str).str.strip()
    em["EquipmentNo"] = em["EquipmentNo"].astype(str).str.strip()
    wl = wl.merge(em, left_on="Suggested EquipmentNo", right_on="EquipmentNo", how="left")

    cls_rows = wl["CatalogProfile"].apply(lambda c: resolve_class(c, lookup))
    wl["Class"] = [r[0] for r in cls_rows]
    wl["ClassLevel"] = [r[1] for r in cls_rows]
    wl["ClassDescription"] = [r[2] for r in cls_rows]
    wl["TaskListGroup"] = wl["Class"].apply(lambda c: f"{PLANT_CODE}{c}" if c else None)

    # Sheet 1: equipment -> class map
    sheet1 = wl[[
        "WorklistID", "Suggested EquipmentNo", "ObjectDescription", "CatalogProfile",
        "Class", "ClassDescription", "TaskListGroup", "Discipline", "Job Description",
    ]].rename(columns={"Suggested EquipmentNo": "EquipmentNo"})

    # ----- Join to tasklist to get all task lines per equipment
    tl["Task"] = tl["Task"].astype(str).str.strip()
    joined = tl.merge(
        wl[["WorklistID", "Suggested EquipmentNo", "TaskListGroup", "Class", "ClassDescription"]],
        on="WorklistID", how="inner",
    ).dropna(subset=["TaskListGroup", "Task"])
    print(f"Joined {len(joined)} task rows across {joined['TaskListGroup'].nunique()} classes")

    # ----- Per class, call the agent
    job_lib_rows = []     # Sheet 2
    counter_eq_rows = []  # Sheet 3

    for (group, cls_code, cls_desc), block in sorted(
        joined.groupby(["TaskListGroup", "Class", "ClassDescription"], dropna=False),
        key=lambda x: x[0][0],
    ):
        # Build per-equipment task sequences
        equipments = []
        eq_lookup = {}  # id -> equipment_no (for sheet 3)
        for eid, (eq_no, erows) in enumerate(
            block.sort_values(["Suggested EquipmentNo", "TaskNo"])
                 .groupby("Suggested EquipmentNo", sort=True), start=1,
        ):
            tasks = erows["Task"].tolist()
            durs = [float(d) if pd.notna(d) else 0.0 for d in erows["Dur (Hrs)"].tolist()]
            equipments.append({"id": eid, "equipment_no": str(eq_no),
                               "tasks": tasks, "durations_hrs": durs})
            eq_lookup[eid] = str(eq_no)

        print(f"[{group}] {cls_desc}: {len(equipments)} equipments, "
              f"{sum(len(e['tasks']) for e in equipments)} task lines -> clustering...")
        resp = call_llm(cls_code, cls_desc,
                        {"equipment_class": f"{cls_code} - {cls_desc}", "equipments": equipments})

        counters = resp.get("counters") or [{
            "header": cls_desc, "operations": [], "equipment_ids": list(eq_lookup),
        }]

        # Emit Sheet 2 (Job Library) and Sheet 3 (counter -> equipments)
        for cnt_idx, cnt in enumerate(counters, start=1):
            ctr = f"{cnt_idx:02d}"
            header = str(cnt.get("header", "")).strip() or f"{cls_desc} #{ctr}"
            ops = cnt.get("operations", []) or []
            total_dur = 0.0

            # Group header row
            job_lib_rows.append({
                "TaskListGroup": group, "GroupCounter": ctr, "Header": header,
                "Plant": PLANT, "ItemNo": "", "OperationDescription": "",
                "Duration (Hrs)": "", "Unit": "", "RowType": "HEADER",
            })
            for op_idx, op in enumerate(ops, start=1):
                dur = float(op.get("duration_hrs") or 0)
                total_dur += dur
                job_lib_rows.append({
                    "TaskListGroup": group, "GroupCounter": ctr, "Header": "",
                    "Plant": PLANT, "ItemNo": f"{op_idx*10:04d}",
                    "OperationDescription": str(op.get("description", "")).strip(),
                    "Duration (Hrs)": round(dur, 2), "Unit": "HR", "RowType": "OPERATION",
                })
            # Totals row
            job_lib_rows.append({
                "TaskListGroup": group, "GroupCounter": ctr, "Header": "",
                "Plant": "", "ItemNo": "", "OperationDescription": "TOTAL",
                "Duration (Hrs)": round(total_dur, 2), "Unit": "HR", "RowType": "TOTAL",
            })

            for eid in cnt.get("equipment_ids", []):
                counter_eq_rows.append({
                    "TaskListGroup": group, "GroupCounter": ctr, "Header": header,
                    "EquipmentNo": eq_lookup.get(eid, f"<unknown id {eid}>"),
                })

    sheet2 = pd.DataFrame(job_lib_rows)
    sheet3 = pd.DataFrame(counter_eq_rows)

    # ----- Write workbook
    with pd.ExcelWriter(OUTPUT, engine="openpyxl") as xw:
        sheet1.to_excel(xw, sheet_name="EquipmentClassMap", index=False)
        sheet2.drop(columns=["RowType"]).to_excel(xw, sheet_name="JobLibrary", index=False)
        sheet3.to_excel(xw, sheet_name="CounterEquipments", index=False)

        # Style the Job Library sheet
        ws = xw.sheets["JobLibrary"]
        hdr_fill = PatternFill("solid", fgColor="305496")
        hdr_font = Font(bold=True, color="FFFFFF")
        tot_fill = PatternFill("solid", fgColor="FFF2CC")
        tot_font = Font(bold=True)
        for i, row_type in enumerate(sheet2["RowType"], start=2):
            if row_type == "HEADER":
                for c in range(1, ws.max_column + 1):
                    ws.cell(i, c).fill = hdr_fill
                    ws.cell(i, c).font = hdr_font
            elif row_type == "TOTAL":
                for c in range(1, ws.max_column + 1):
                    ws.cell(i, c).fill = tot_fill
                    ws.cell(i, c).font = tot_font
        for c in range(1, ws.max_column + 1):
            ws.column_dimensions[ws.cell(1, c).column_letter].width = 22

    print(f"\nWrote {OUTPUT}")
    print(f"Sheet1: {len(sheet1)} equipment rows")
    print(f"Sheet2: {sheet2['RowType'].value_counts().to_dict()}")
    print(f"Sheet3: {len(sheet3)} equipment-to-counter rows "
          f"across {sheet3['GroupCounter'].nunique()} counters")


if __name__ == "__main__":
    try:
        build()
    finally:
        langfuse.flush()
