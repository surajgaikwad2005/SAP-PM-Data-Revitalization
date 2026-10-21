"""Build a 2024 Job Library (IA05-style) from the cleaned worklist + Tasklist_2024.

Per equipment class, cluster similar tasks into Group Counters using the LLM
(Claude Sonnet via Azure Foundry, reused from agents.py).

Output structure mirrors IA05_Planning Norms Electrical.xlsx:
    - Group header rows   (TL Ind=A, Item Flag=H) per counter
    - Operation rows      (Item Flag=I) per task in the cluster
"""

import pandas as pd
from pathlib import Path
from agents import call_llm

BASE = Path(__file__).resolve().parents[1]
ENRICHED = BASE / "Processed Data" / "TaskList_Groups_Bokor.xlsx"
TASKLIST = BASE / "Merged Sheet" / "Tasklist_2024.xlsx"
OUTPUT = BASE / "Processed Data" / "Job_Library_Bokor_2024.xlsx"

PLANT = "M010"               # SAP WERKS for Bokor
WORK_CENTER_BY_DISC = {
    "Electrical": "PMO-ELE",
    "Instrument & Control": "PMO-INS",
    "Static": "PMO-MEC",
    "Rotating": "PMO-MEC",
    "Piping": "PMO-MEC",
    "Construction": "PMO-CON",
    "Mechanical": "PMO-MEC",
}
DEFAULT_WORK_CENTER = "PMO-GEN"


# ---------------------------------------------------------------------------
def cluster_tasks(class_code, class_desc, tasks):
    """Ask the LLM to group similar tasks. tasks = list[(id, text, discipline)].
    Returns list[dict(header=str, task_ids=list[int])]."""
    if len(tasks) <= 1:
        return [{"header": tasks[0][1] if tasks else "", "task_ids": [tasks[0][0]] if tasks else []}]

    system = (
        "You are a SAP PM maintenance planner building a Job Library. "
        "Given a list of individual maintenance task descriptions for one equipment class, "
        "cluster them into groups of SIMILAR OPERATIONS. Tasks in the same cluster describe "
        "essentially the same work step (same intent, same procedure), even if wording differs. "
        "Give each cluster a short, canonical header (<=60 chars, title case). "
        "Return STRICT JSON: {\"clusters\":[{\"header\":\"...\",\"task_ids\":[1,2,...]},...]}. "
        "Every input task_id must appear in exactly one cluster."
    )
    payload = {
        "equipment_class": f"{class_code} - {class_desc}",
        "tasks": [{"id": tid, "text": txt, "discipline": disc} for tid, txt, disc in tasks],
    }
    resp = call_llm(system, payload, max_tokens=2000)
    if "_error" in resp or "clusters" not in resp:
        print(f"  [LLM failed for {class_code}; falling back to one-cluster-per-task]")
        return [{"header": t[1][:60], "task_ids": [t[0]]} for t in tasks]
    return resp["clusters"]


def work_center(discipline):
    return WORK_CENTER_BY_DISC.get(str(discipline).strip(), DEFAULT_WORK_CENTER)


# ---------------------------------------------------------------------------
# Load data
enriched = pd.read_excel(ENRICHED, sheet_name="Worklist_Enriched")
enriched = enriched.dropna(subset=["TaskListGroup"])[
    ["WorklistID", "TaskListGroup", "Class", "ClassDescription",
     "Suggested EquipmentNo", "Discipline"]
]

tl = pd.read_excel(TASKLIST)
tl = tl[["WorklistID", "EquipmentNo", "TaskNo", "Task", "Dur (Hrs)", "Discipline"]].dropna(subset=["Task"])
tl["Task"] = tl["Task"].astype(str).str.strip()

joined = tl.merge(enriched, on="WorklistID", how="inner", suffixes=("", "_wl"))
joined = joined.dropna(subset=["TaskListGroup", "Task"])
print(f"Joined {len(joined)} task rows across {joined['TaskListGroup'].nunique()} classes")

# ---------------------------------------------------------------------------
# Build the Job Library per class
rows = []
for (group, cls_code, cls_desc), block in sorted(
    joined.groupby(["TaskListGroup", "Class", "ClassDescription"], dropna=False),
    key=lambda x: x[0][0],
):
    # Dedup tasks within the class (case-insensitive) so clustering is efficient
    dedup = block.drop_duplicates(subset=[block["Task"].str.lower().name])
    uniq = (
        block.assign(_key=block["Task"].str.lower().str.strip())
             .groupby("_key", as_index=False)
             .agg(Task=("Task", "first"),
                  Dur=("Dur (Hrs)", "mean"),
                  Discipline=("Discipline", lambda s: s.mode().iat[0] if len(s.mode()) else ""))
    )
    uniq = uniq.reset_index(drop=True)
    print(f"\n[{group}] {cls_desc}: {len(uniq)} unique tasks -> clustering...")

    tasks = list(zip(uniq.index.tolist(), uniq["Task"].tolist(), uniq["Discipline"].tolist()))
    clusters = cluster_tasks(cls_code, cls_desc, tasks)

    # Emit IA05-shaped rows
    for cnt_idx, cluster in enumerate(clusters, start=1):
        counter = f"{cnt_idx:02d}"
        header = str(cluster.get("header", "")).strip() or f"{cls_desc} Group {counter}"
        task_ids = cluster.get("task_ids", [])

        # Header row (TL Ind=A, Item Flag=H)
        rows.append({
            "TCODE": "IA05", "TL Ind": "A", "GenTL (Group)": group,
            "Grp Counter": counter, "Header Description": header,
            "Item Flag": "H", "Item No": "", "Plant": PLANT,
            "Work Center": "", "Operation Description": "",
            "Duration (Hrs)": "", "Unit": "",
        })

        # Operation rows (Item Flag=I) under this counter
        for op_idx, tid in enumerate(task_ids, start=1):
            if tid >= len(uniq):
                continue
            trow = uniq.iloc[tid]
            rows.append({
                "TCODE": "IA05", "TL Ind": "", "GenTL (Group)": "",
                "Grp Counter": counter, "Header Description": "",
                "Item Flag": "I", "Item No": f"{op_idx*10:04d}", "Plant": PLANT,
                "Work Center": work_center(trow["Discipline"]),
                "Operation Description": trow["Task"],
                "Duration (Hrs)": round(float(trow["Dur"]) if pd.notna(trow["Dur"]) else 0, 2),
                "Unit": "HR",
            })

# ---------------------------------------------------------------------------
# Write output with styled headers
from openpyxl.styles import PatternFill, Font

job_lib = pd.DataFrame(rows)
with pd.ExcelWriter(OUTPUT, engine="openpyxl") as xw:
    job_lib.to_excel(xw, sheet_name="JobLibrary", index=False)
    ws = xw.sheets["JobLibrary"]
    fill = PatternFill("solid", fgColor="305496")
    font = Font(bold=True, color="FFFFFF")
    for r in range(2, ws.max_row + 1):
        if (ws.cell(r, 6).value or "") == "H":   # Item Flag column
            for c in range(1, ws.max_column + 1):
                ws.cell(r, c).fill = fill
                ws.cell(r, c).font = font
    for c in range(1, ws.max_column + 1):
        ws.column_dimensions[ws.cell(1, c).column_letter].width = 22

print(f"\nWrote {OUTPUT}")
n_counters = job_lib[job_lib["Item Flag"] == "H"].shape[0]
n_ops = job_lib[job_lib["Item Flag"] == "I"].shape[0]
print(f"{n_counters} group counters across {job_lib['GenTL (Group)'].replace('', None).dropna().nunique()} "
      f"tasklist groups, {n_ops} operation lines")
